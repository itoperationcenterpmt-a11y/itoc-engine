"""Sangfor IAG Report Center exports (.xlsx, one file per report) into one summary per site and period.
Only the network figures are read: bandwidth per application category, traffic ranking with top users and applications, users per category, rejected requests, and the daily
connection-sharing and mobile-endpoint trends. The per-person reports (Work Efficiency Assessment, Employee Turnover Risk) are skipped on purpose.
Volumes in the exports are in gigabits (Gb = bits); they are stored as gigabytes. Rates stay in Mb/s."""
import re
from datetime import date
from io import BytesIO

_UNIT = {"b": 1.0, "kb": 1e3, "mb": 1e6, "gb": 1e9, "tb": 1e12}
_VOL = re.compile(r"^\s*([\d.,]+)\s*([KMGT]?b)\s*$", re.I)
_RATE = re.compile(r"^\s*([\d.,]+)\s*([KMGT]?b)/s\s*$", re.I)
_ITEM = re.compile(r"^(.*?):\s*([\d.,]+\s*[KMGT]?b(?:/s)?)\s*(?:\(([\d.]+)%\))?\s*$", re.I)
SKIPPED = ("Work Efficiency Assessment", "Employee Turnover Risk")


def gb(v) -> float | None:
    """'55663.92Gb' (gigabits) -> gigabytes."""
    m = _VOL.match(str(v or ""))
    return round(float(m.group(1).replace(",", "")) * _UNIT[m.group(2).lower()] / 8 / 1e9, 3) if m else None


def mbps(v) -> float | None:
    m = _RATE.match(str(v or ""))
    return round(float(m.group(1).replace(",", "")) * _UNIT[m.group(2).lower()] / 1e6, 2) if m else None


def pct(v) -> float | None:
    try:
        return float(str(v).replace("%", "").replace(",", ".").strip())
    except ValueError:
        return None


def num(v) -> int | None:
    try:
        return int(float(str(v).replace(",", "").strip()))
    except ValueError:
        return None


def items(cell, rate: bool = False) -> list:
    """'10.181.3.3:105909.33Gb \\n10.181.2.2:1807Gb \\nOthers:32880Gb' -> [[name, value, percent]]."""
    out = []
    for ln in str(cell or "").split("\n"):
        m = _ITEM.match(ln.strip())
        if m:
            out.append([m.group(1).strip(), mbps(m.group(2)) if rate else gb(m.group(2)), pct(m.group(3)) if m.group(3) else None])
    return out


def _rows(ws) -> tuple[dict, list, list]:
    """(options shown above 'Results', header, data rows) of one export sheet."""
    opts, head, data, seen = {}, None, [], False
    for r in ws.iter_rows(values_only=True):
        v = [c for c in r]
        if not seen:
            if v and v[0] == "Results":
                seen = True
            elif v and v[0] and len(v) > 1 and v[1] is not None:
                opts[str(v[0])] = str(v[1])
            continue
        if head is None:
            head = [str(c or "") for c in v]
        elif any(c is not None and str(c).strip() for c in v):
            data.append(v)
    return opts, head or [], data


def _period(opts) -> tuple[date, date] | None:
    m = re.search(r"(\d{4})-(\d\d)-(\d\d).*?(\d{4})-(\d\d)-(\d\d)", opts.get("Date Range", ""))
    return (date(int(m.group(1)), int(m.group(2)), int(m.group(3))), date(int(m.group(4)), int(m.group(5)), int(m.group(6)))) if m else None


def parse(files: list[tuple[str, bytes]]) -> dict:
    """files = [(name, bytes)] -> {"from", "to", "data": {...}, "read": [...], "skipped": [...]}"""
    import openpyxl
    data, read, skipped, per = {}, [], [], []
    for name, raw in files:
        try:
            ws = openpyxl.load_workbook(BytesIO(raw), read_only=True, data_only=True).worksheets[0]
        except Exception:
            skipped.append(f"{name}: not a readable Excel file")
            continue
        first = next(ws.iter_rows(min_row=1, max_row=1, values_only=True), (None,))
        title = str(first[0] or "").strip()
        if title in SKIPPED:
            skipped.append(f"{title}: per-person report, not imported")
            continue
        opts, head, rows = _rows(ws)
        col = lambda n: head.index(n) if n in head else None
        p = _period(opts)
        if p:
            per.append(p)
        if title == "Application Bandwidth Distribution":
            first_col = head[1] if len(head) > 1 else ""
            key = "cat" if first_col == "App Category" else "dir"
            out = []
            for r in rows:
                d = dict(zip(head, r))
                out.append({"k": str(d.get(first_col) or ""), "max": mbps(d.get("Max Bidirectional")), "avg": mbps(d.get("Avg Bidirectional")), "out": gb(d.get("Outbound")), "in": gb(d.get("Inbound")),
                            "bi": gb(d.get("Bidirectional")), "pct": pct(d.get("Total Bandwidth Percent") or d.get("Bandwidth Percent")), "users": items(d.get("Users"), rate=True)})
            data[key] = out
        elif title == "Application Traffic Ranking":
            data["rank"] = [{"k": str(d.get("App Category") or ""), "users": items(d.get("Username")), "apps": items(d.get("Application")), "out": gb(d.get("Outbound")), "in": gb(d.get("Inbound")),
                             "bi": gb(d.get("Bidirectional")), "pct": pct(d.get("Bidirectional Percent"))} for d in (dict(zip(head, r)) for r in rows)]
        elif title == "Application Access Ranking":
            data["access"] = [{"k": str(d.get("App Category") or ""), "users": num(d.get("Users")), "pct": pct(d.get("Percent"))} for d in (dict(zip(head, r)) for r in rows)]
        elif title == "Application Access Trend":
            data["reject"] = [{"k": str(d.get("Application") or ""), "reject": num(d.get("Reject")), "count": num(d.get("Access Count"))} for d in (dict(zip(head, r)) for r in rows)]
        elif title in ("Connection Sharing Trend", "Mobile Endpoint Access Trend"):
            key = "sharing" if title.startswith("Connection") else "mobile"
            last = head[-1] if head else ""
            data[key] = [[re.sub(r"^(\d{4})(\d\d)(\d\d)$", r"\1-\2-\3", str(d.get("Date Range") or "")), num(d.get(last))] for d in (dict(zip(head, r)) for r in rows)]
        else:
            skipped.append(f"{title or name}: not a network report")
            continue
        read.append(title)
    if not data:
        raise ValueError("No Sangfor Report Center export was recognised in these files.")
    if not per:
        raise ValueError("The period (Date Range) could not be read from the files.")
    return {"from": min(p[0] for p in per), "to": max(p[1] for p in per), "data": data, "read": read, "skipped": skipped}
