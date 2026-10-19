"""Pembaca workbook (openpyxl untuk .xlsx, xlrd untuk .xls ekspor lama) + parser tanggal yang sama dengan versi Excel/JS."""
import re
from datetime import datetime, date, timedelta
from io import BytesIO
from openpyxl import load_workbook

MON = {m: i + 1 for i, m in enumerate(["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"])}
_R1 = re.compile(r"^(?:[A-Za-z]{3},\s*)?([A-Za-z]{3})[a-z]*\.?\s+(\d{1,2}),\s*(\d{4})(?:\s+(\d{1,2}):(\d{2})(?::\d{2})?\s*([AaPp][Mm])?)?$")
_R2 = re.compile(r"^(\d{4})-(\d{2})-(\d{2})(?:[ T](\d{2}):(\d{2}))?")
_R3 = re.compile(r"^(\d{1,2})/(\d{1,2})/(\d{4})(?:\s+(\d{1,2}):(\d{2})(?::\d{2})?\s*([AaPp][Mm])?)?")


def s(v) -> str:
    return "" if v is None else str(v).strip()


def num(v) -> float:
    if isinstance(v, (int, float)):
        return float(v)
    try:
        return float(s(v).replace(",", ""))
    except ValueError:
        return 0.0


def _hour(h, ap):
    h = int(h) if h else 0
    if ap:
        h = h % 12 + (12 if ap.lower().startswith("p") else 0)
    return h


def pdt(v) -> datetime | None:
    """Tanggal ekspor → datetime naif (jam lokal). Format yang tidak dikenal/invalid → None (baris dilewati)."""
    if v is None or v == "":
        return None
    if isinstance(v, datetime):
        return v.replace(tzinfo=None)
    if isinstance(v, date):
        return datetime.combine(v, datetime.min.time())
    if isinstance(v, (int, float)):
        return datetime(1899, 12, 30) + timedelta(days=float(v))
    t = s(v)
    try:
        if m := _R1.match(t):
            mo = MON.get(m[1].lower())
            return None if mo is None else datetime(int(m[3]), mo, int(m[2]), _hour(m[4], m[6]), int(m[5] or 0))
        if m := _R2.match(t):
            return datetime(int(m[1]), int(m[2]), int(m[3]), int(m[4] or 0), int(m[5] or 0))
        if m := _R3.match(t):
            return datetime(int(m[3]), int(m[1]), int(m[2]), _hour(m[4], m[6]), int(m[5] or 0))
    except ValueError:
        return None
    return None


class Workbook:
    """Membuka .xlsx (zip) maupun .xls biner (OLE2, ekspor ManageEngine lama) dari isi file, bukan dari ekstensi."""

    def __init__(self, data: bytes):
        self._cache = {}
        self.xls = data[:4] == b"\xd0\xcf\x11\xe0"
        if self.xls:
            import xlrd
            self.wb = xlrd.open_workbook(file_contents=data, on_demand=True)
        else:
            self.wb = load_workbook(BytesIO(data), read_only=True, data_only=True)

    @property
    def names(self):
        return self.wb.sheet_names() if self.xls else self.wb.sheetnames

    def has(self, n):
        return n in self.names

    def rows(self, n) -> list[list]:
        if n not in self._cache:
            if not self.has(n):
                self._cache[n] = []
            elif self.xls:
                sh = self.wb.sheet_by_name(n)
                self._cache[n] = [sh.row_values(i) for i in range(sh.nrows)]
            else:
                self._cache[n] = [list(r) for r in self.wb[n].iter_rows(values_only=True)]
        return self._cache[n]

    def cell(self, n, addr):
        """Alamat A1. Membaca dari cache baris (read_only tidak mendukung akses acak)."""
        m = re.match(r"([A-Z]+)(\d+)", addr)
        col = 0
        for ch in m[1]:
            col = col * 26 + ord(ch) - 64
        r = self.rows(n)
        i = int(m[2]) - 1
        return r[i][col - 1] if i < len(r) and col - 1 < len(r[i]) else None
