"""Penyusun file laporan dari model `report.build`: Excel (openpyxl), PDF (reportlab), PowerPoint (python-pptx, grafik asli).
Pustaka diimpor saat dipakai supaya aplikasi tetap jalan walau salah satu belum terpasang."""
from io import BytesIO
from pathlib import Path
from xml.sax.saxutils import escape
from .report import fmt

NAVY, BLUE, LINE, MUTED = "0F3D7A", "1C5CAB", "DDE2E7", "6F7B87"
PAL = ["2A78D6", "EB8A34", "0CA30C", "D03B3B", "7A5AF5", "14B8A6", "F2B84B", "6B7280"]
XL_FMT = {"n": "#,##0", "f0": "#,##0", "f1": "#,##0.0", "f2": "#,##0.00", "p0": "0%", "p1": "0.0%", "p2": "0.00%"}
MIME = {"xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", "pdf": "application/pdf",
        "pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation"}


# ---------------------------------------------------------------- Excel
def to_xlsx(m: dict) -> bytes:
    from openpyxl import Workbook
    from openpyxl.chart import AreaChart, BarChart, DoughnutChart, LineChart, Reference
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter
    head_fill, thin = PatternFill("solid", fgColor=NAVY), Side(style="thin", color=LINE)
    wb = Workbook()
    wb.remove(wb.active)
    for s in m["sections"]:
        ws = wb.create_sheet(s["title"][:31])
        ws.sheet_view.showGridLines = False
        width = {}

        def put(r, c, v, kind="t", bold=False, color=None, fill=None, wrap=False):
            cell = ws.cell(row=r, column=c, value=None if v is None else v)
            if kind in XL_FMT and isinstance(v, (int, float)):
                cell.number_format = XL_FMT[kind]
            cell.font = Font(bold=bold, color=color or "10171E", size=10)
            if fill:
                cell.fill = fill
            cell.alignment = Alignment(vertical="center", wrap_text=wrap, horizontal="right" if kind != "t" and isinstance(v, (int, float)) else "left")
            cell.border = Border(bottom=thin)
            width[c] = min(60, max(width.get(c, 8), len(fmt(v, kind)) + 2))
            return cell

        ws["A1"] = m["title"]
        ws["A1"].font = Font(bold=True, size=16, color=NAVY)
        ws["A2"] = f"{m['subtitle']} · {s['title']} · {m['period']}"
        ws["A2"].font = Font(bold=True, size=11)
        ws["A3"] = " · ".join(m["filters"])
        ws["A4"] = f"Generated {m['generated']} by {m['by']}"
        for a in ("A3", "A4"):
            ws[a].font = Font(size=9, color=MUTED)
        r = 6
        for n in s["notes"]:
            ws.cell(row=r, column=1, value=n).font = Font(italic=True, color=MUTED)
            r += 1
        if s["kpis"]:
            for c, h in enumerate(["Metrik", "Nilai", "Keterangan"], 1):
                put(r, c, h, bold=True, color="FFFFFF", fill=head_fill)
            r += 1
            for k in s["kpis"]:
                put(r, 1, k["label"])
                put(r, 2, k["value"], k["kind"], bold=True)
                put(r, 3, k["sub"])
                r += 1
            r += 1
        # grafik: data sumber di kolom AD ke kanan; grafik tersusun 2 kolom di bawah KPI
        gcol, grow, cr = 30, 6, r
        for idx, ch in enumerate(s["charts"]):
            put(grow, gcol, ch["title"], bold=True)
            hdr = grow + 1
            put(hdr, gcol, "Category", bold=True, color="FFFFFF", fill=head_fill)
            for j, se in enumerate(ch["series"]):
                put(hdr, gcol + 1 + j, se["name"], bold=True, color="FFFFFF", fill=head_fill)
            nlab, nser = len(ch["labels"]), len(ch["series"])
            for i, lab in enumerate(ch["labels"]):
                put(hdr + 1 + i, gcol, lab)
                for j, se in enumerate(ch["series"]):
                    put(hdr + 1 + i, gcol + 1 + j, se["values"][i], "f1" if isinstance(se["values"][i], float) else "n")
            kind = ch["kind"]
            if kind == "donut":
                bc = DoughnutChart()
                bc.holeSize = 55
            elif kind in ("line", "prtg"):
                bc = LineChart()
            elif kind == "area":
                bc = AreaChart()
            else:
                bc = BarChart()
                bc.type = "bar" if kind == "hbar" else "col"
                if kind == "stack":
                    bc.grouping, bc.overlap = "stacked", 100
                if kind == "hbar":
                    bc.x_axis.scaling.orientation = "maxMin"
            bc.title, bc.height, bc.width = ch["title"], 7.4, 14.2
            bc.add_data(Reference(ws, min_col=gcol + 1, max_col=gcol + nser, min_row=hdr, max_row=hdr + nlab), titles_from_data=True)
            bc.set_categories(Reference(ws, min_col=gcol, min_row=hdr + 1, max_row=hdr + nlab))
            if kind in ("area", "bar", "stack", "hbar"):
                for j, se in enumerate(ch["series"]):
                    if se.get("color") and j < len(bc.series):
                        bc.series[j].graphicalProperties.solidFill = se["color"]
            if nser == 1 and kind != "donut":
                bc.legend = None
            ws.add_chart(bc, f"{'A' if idx % 2 == 0 else 'G'}{cr + (idx // 2) * 16}")
            grow = hdr + nlab + 3
        if s["charts"]:
            r = cr + ((len(s["charts"]) + 1) // 2) * 16 + 1
        for t in s["tables"]:
            ws.cell(row=r, column=1, value=t["title"]).font = Font(bold=True, size=12, color=BLUE)
            r += 1
            for c, h in enumerate(t["heads"], 1):
                put(r, c, h, bold=True, color="FFFFFF", fill=head_fill, wrap=True)
            r += 1
            for row in t["rows"]:
                for c, v in enumerate(row, 1):
                    put(r, c, v, t["kinds"][c - 1])
                r += 1
            r += 1
        for c, w in width.items():
            ws.column_dimensions[get_column_letter(c)].width = w if c < 12 else max(w, 10)
        ws.column_dimensions["A"].width = max(ws.column_dimensions["A"].width or 0, 34)
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


# ---------------------------------------------------------------- PDF
def _pdf_area(ch, w, h):
    """Traffic graph in the PRTG style: filled areas under the lines, light grid, few date labels, legend underneath."""
    import math
    from reportlab.graphics.shapes import Drawing, Line, Polygon, PolyLine, Rect, String
    from reportlab.lib import colors
    col = lambda c, a=1.0: (lambda x: colors.Color(x.red, x.green, x.blue, alpha=a))(colors.HexColor("#" + c))
    d = Drawing(w, h)
    d.add(String(4, h - 10, ch["title"], fontName="Helvetica-Bold", fontSize=8.5, fillColor=col(NAVY)))
    labels, series, n = ch["labels"], ch["series"], len(ch["labels"])
    x0, y0, pw, ph = 40, 34, w - 52, h - 58
    vmax = max([v or 0 for s in series for v in s["values"]] + [0.0001])
    raw = vmax / 4
    mag = 10 ** math.floor(math.log10(raw))
    step = next(m * mag for m in (1, 2, 2.5, 5, 10) if m * mag >= raw)
    top = step * math.ceil(vmax / step)
    nice = lambda v: ("%.2f" % v).rstrip("0").rstrip(".") if v < 10 else "%d" % round(v)
    for k in range(int(round(top / step)) + 1):
        y = y0 + ph * k * step / top
        d.add(Line(x0, y, x0 + pw, y, strokeColor=col("E1E7EF"), strokeWidth=.5))
        d.add(String(x0 - 4, y - 2.5, nice(k * step), fontName="Helvetica", fontSize=6.5, textAnchor="end", fillColor=col("5B6B7F")))
    xs = [x0 + (pw * i / (n - 1) if n > 1 else pw / 2) for i in range(n)]
    for j, se in enumerate(series):
        c = se.get("color") or PAL[j % len(PAL)]
        pts = [(xs[i], y0 + ph * (v or 0) / top) for i, v in enumerate(se["values"])]
        flat = [p for xy in pts for p in xy]
        d.add(Polygon([xs[0], y0] + flat + [xs[-1], y0], fillColor=col(c, .38), strokeColor=None))
        d.add(PolyLine(flat, strokeColor=col(c), strokeWidth=1.1))
    d.add(Line(x0, y0, x0 + pw, y0, strokeColor=col("9AA7B8"), strokeWidth=.7))
    idx = sorted({round(i * (n - 1) / 6) for i in range(7)}) if n > 7 else range(n)
    for i in idx:
        d.add(String(xs[i], y0 - 10, labels[i], fontName="Helvetica", fontSize=6.5, textAnchor="middle", fillColor=col("5B6B7F")))
    lx = x0
    for j, se in enumerate(series):
        d.add(Rect(lx, 6, 8, 6, fillColor=col(se.get("color") or PAL[j % len(PAL)], .6), strokeColor=None))
        d.add(String(lx + 11, 6, se["name"], fontName="Helvetica", fontSize=7, fillColor=col("33445C")))
        lx += 24 + len(se["name"]) * 4
    return d


def _pdf_prtg(ch, w, h):
    """One link in the PRTG style: traffic lines with the highest and lowest point marked, downtime as red vertical lines in the same graph, caption and footer."""
    import math
    from reportlab.graphics.shapes import Circle, Drawing, Line, PolyLine, Rect, String
    from reportlab.lib import colors
    col = lambda c: colors.HexColor("#" + c)
    d = Drawing(w, h)
    labels, series, down, n = ch["labels"], ch["series"], ch["down"], len(ch["labels"])
    d.add(String(4, h - 10, ch["title"], fontName="Helvetica-Bold", fontSize=9, fillColor=col(NAVY)))
    if ch.get("sub"):
        d.add(String(w - 4, h - 10, ch["sub"], fontName="Helvetica", fontSize=7, textAnchor="end", fillColor=col("5B6B7F")))
    x0, y0, pw, ph = 40, 44, w - 78, h - 68
    vmax = max([v or 0 for s in series for v in s["values"]] + [0.0001])
    raw = vmax / 4
    mag = 10 ** math.floor(math.log10(raw))
    step = next(m * mag for m in (1, 2, 2.5, 5, 10) if m * mag >= raw)
    top = step * math.ceil(vmax / step)
    nice = lambda v: ("%.2f" % v).rstrip("0").rstrip(".") if v < 10 else "%d" % round(v)
    for k in range(int(round(top / step)) + 1):
        y = y0 + ph * k * step / top
        d.add(Line(x0, y, x0 + pw, y, strokeColor=col("DDDDDD"), strokeWidth=.5))
        d.add(String(x0 - 4, y - 2.5, nice(k * step), fontName="Helvetica", fontSize=6.5, textAnchor="end", fillColor=col("5B6B7F")))
    for k in range(5):
        d.add(String(x0 + pw + 4, y0 + ph * k / 4 - 2.5, f"{k * 25}%", fontName="Helvetica", fontSize=6.5, fillColor=col("D6558F")))
    d.add(String(x0 - 4, y0 + ph + 5, ch.get("unit", ""), fontName="Helvetica", fontSize=6.5, textAnchor="end", fillColor=col("5B6B7F")))
    xs = [x0 + (pw * i / (n - 1) if n > 1 else pw / 2) for i in range(n)]
    for i, v in enumerate(down):
        if v and v > 0:
            bw = max(pw / max(n, 1), .7)
            d.add(Rect(xs[i] - bw / 2, y0, bw, ph, fillColor=colors.Color(.85, .33, .33, alpha=.13), strokeColor=None))
            d.add(Rect(xs[i] - bw / 2, y0, bw, ph * min(max(v, 6), 100) / 100, fillColor=colors.Color(.85, .33, .33, alpha=.45), strokeColor=None))
    for j, se in reversed(list(enumerate(series))):
        pts = [p for i, v in enumerate(se["values"]) for p in (xs[i], y0 + ph * (v or 0) / top)]
        d.add(PolyLine(pts, strokeColor=col(se["color"]), strokeWidth=1.0 if j else 1.3))
    tot = series[0]["values"]
    if tot:
        up = [v for v in tot if v and v > 0] or tot
        for fn, tag in ((max, "Max"), (min, "Min")):
            i = tot.index(fn(up))
            x, y = xs[i], y0 + ph * tot[i] / top
            d.add(Circle(x, y, 2.2, fillColor=col(series[0]["color"]), strokeColor=colors.white, strokeWidth=.6))
            d.add(String(min(max(x, x0 + 30), x0 + pw - 30), y + 4, f"{tag}: {nice(tot[i])} {ch.get('unit', '')}", fontName="Helvetica-Bold", fontSize=6.5, textAnchor="middle", fillColor=col("33445C")))
    d.add(Line(x0, y0, x0 + pw, y0, strokeColor=col("9AA7B8"), strokeWidth=.7))
    for i in (sorted({round(k * (n - 1) / 7) for k in range(8)}) if n > 8 else range(n)):
        d.add(String(xs[i], y0 - 9, labels[i], fontName="Helvetica", fontSize=6, textAnchor="middle", fillColor=col("5B6B7F")))
    lx = x0
    for se in series:
        d.add(Line(lx, 24, lx + 10, 24, strokeColor=col(se["color"]), strokeWidth=1.6))
        d.add(String(lx + 13, 21.5, se["name"], fontName="Helvetica", fontSize=7, fillColor=col("33445C")))
        lx += 30 + len(se["name"]) * 3.6
    d.add(Rect(lx, 20, 6, 8, fillColor=colors.Color(.85, .33, .33, alpha=.5), strokeColor=None))
    d.add(String(lx + 10, 21.5, "Downtime (right axis: share of the hour or day that was down)", fontName="Helvetica", fontSize=7, fillColor=col("33445C")))
    d.add(String(x0, 6, f"TERA · {ch.get('res', '')} · ID {ch.get('sid', '')}", fontName="Helvetica", fontSize=6, fillColor=col("8A96A6")))
    return d


def _pdf_chart(ch, w, h):
    """Satu grafik reportlab (batang, tumpuk, horizontal, donut, garis) dalam Drawing berukuran w x h."""
    from reportlab.graphics.charts.barcharts import HorizontalBarChart, VerticalBarChart
    from reportlab.graphics.charts.linecharts import HorizontalLineChart
    from reportlab.graphics.charts.piecharts import Pie
    from reportlab.graphics.shapes import Drawing, Rect, String
    from reportlab.lib import colors
    hx = lambda c: colors.HexColor("#" + c)
    d = Drawing(w, h)
    d.add(String(4, h - 10, ch["title"], fontName="Helvetica-Bold", fontSize=8.5, fillColor=hx(NAVY)))
    kind, series = ch["kind"], ch["series"]
    if kind == "prtg":
        return _pdf_prtg(ch, w, h)
    if kind == "area":
        return _pdf_area(ch, w, h)
    labels = [(l if len(l) <= 34 else l[:33] + "…") for l in ch["labels"]]
    data = [[v or 0 for v in se["values"]] for se in series]
    if kind == "donut":
        size = h - 40
        p = Pie()
        p.x, p.y, p.width, p.height = 14, 14, size, size
        p.data = data[0]
        p.slices.strokeColor = colors.white
        p.slices.strokeWidth = 1
        for i in range(len(data[0])):
            p.slices[i].fillColor = hx(PAL[i % len(PAL)])
        d.add(p)
        tot = sum(data[0]) or 1
        for i, lab in enumerate(labels[:10]):
            y = h - 34 - i * 13
            d.add(Rect(size + 30, y - 1, 7, 7, fillColor=hx(PAL[i % len(PAL)]), strokeColor=None))
            d.add(String(size + 42, y, f"{lab}: {fmt(data[0][i], 'n')} ({data[0][i] / tot * 100:.0f}%)", fontName="Helvetica", fontSize=7.5))
        return d
    if kind == "hbar":
        b = HorizontalBarChart()
        b.x, b.y, b.width, b.height = 150, 12, w - 165, h - 36
        b.data = [list(reversed(row)) for row in data]
        b.categoryAxis.categoryNames = list(reversed(labels))
        b.categoryAxis.labels.boxAnchor, b.categoryAxis.labels.dx = "e", -3
        b.categoryAxis.labels.fontSize = 6.5
        b.valueAxis.labels.fontSize = 6.5
        b.valueAxis.valueMin = 0
        b.barSpacing = 1
    elif kind == "line":
        b = HorizontalLineChart()
        b.x, b.y, b.width, b.height = 36, 34, w - 50, h - 60
        b.data = data
        b.categoryAxis.categoryNames = labels
        b.categoryAxis.labels.fontSize = 6
        b.valueAxis.labels.fontSize = 6.5
        for i in range(len(data)):
            b.lines[i].strokeColor = hx(PAL[i % len(PAL)])
            b.lines[i].strokeWidth = 1.5
        if len(labels) > 8:
            b.categoryAxis.labels.angle = 50
            b.categoryAxis.labels.boxAnchor = "ne"
            b.categoryAxis.labels.dx, b.categoryAxis.labels.dy = -2, -2
    else:
        b = VerticalBarChart()
        b.x, b.y, b.width, b.height = 36, 34, w - 50, h - 60
        b.data = data
        b.categoryAxis.categoryNames = labels
        b.categoryAxis.labels.fontSize = 6
        b.valueAxis.labels.fontSize = 6.5
        b.valueAxis.valueMin = 0
        b.groupSpacing = 3
        if kind == "stack":
            b.categoryAxis.style = "stacked"
        if len(labels) > 8:
            b.categoryAxis.labels.angle = 50
            b.categoryAxis.labels.boxAnchor = "ne"
            b.categoryAxis.labels.dx, b.categoryAxis.labels.dy = -2, -2
    if kind != "line":
        for j in range(len(data)):
            b.bars[j].fillColor = hx(series[j].get("color") or PAL[j % len(PAL)])
            b.bars[j].strokeColor = None
    d.add(b)
    if len(series) > 1:
        lx = 40
        for i, se in enumerate(series):
            d.add(Rect(lx, 3, 7, 7, fillColor=hx(PAL[i % len(PAL)]), strokeColor=None))
            d.add(String(lx + 9, 3, se["name"], fontName="Helvetica", fontSize=7))
            lx += 20 + len(se["name"]) * 4
    return d


def to_pdf(m: dict) -> bytes:
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib.styles import ParagraphStyle
    from reportlab.lib.units import mm
    from reportlab.platypus import Image, KeepTogether, PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle
    W = landscape(A4)[0] - 24 * mm
    hx = lambda c: colors.HexColor("#" + c)
    st = {"h1": ParagraphStyle("h1", fontName="Helvetica-Bold", fontSize=20, leading=24, textColor=hx(NAVY)),
          "h2": ParagraphStyle("h2", fontName="Helvetica-Bold", fontSize=14, leading=18, textColor=hx(NAVY), spaceBefore=4, spaceAfter=4),
          "h3": ParagraphStyle("h3", fontName="Helvetica-Bold", fontSize=10.5, leading=14, textColor=hx(BLUE), spaceBefore=8, spaceAfter=3, keepWithNext=1),
          "sub": ParagraphStyle("sub", fontName="Helvetica", fontSize=9, leading=12, textColor=hx(MUTED)),
          "lab": ParagraphStyle("lab", fontName="Helvetica", fontSize=7.5, leading=9, textColor=hx(MUTED)),
          "val": ParagraphStyle("val", fontName="Helvetica-Bold", fontSize=17, leading=20, textColor=hx(NAVY)),
          "c": ParagraphStyle("c", fontName="Helvetica", fontSize=7.5, leading=9),
          "ch": ParagraphStyle("ch", fontName="Helvetica-Bold", fontSize=7.5, leading=9, textColor=colors.white),
          "note": ParagraphStyle("note", fontName="Helvetica-Oblique", fontSize=9, leading=12, textColor=hx(MUTED))}
    P = lambda t, s: Paragraph(escape(str(t)), st[s])
    head = [P(m["title"], "h1"), Spacer(1, 3), P(f"{m['subtitle']} · {m['period']}", "h2"), P(" · ".join(m["filters"]), "sub"), P(f"Generated {m['generated']} by {m['by']}", "sub")]
    logo = Path(__file__).resolve().parent.parent / "assets" / "logo.png"
    if logo.exists():
        ht = Table([[Image(str(logo), 17 * mm, 17 * mm), head]], colWidths=[22 * mm, W - 22 * mm])
        ht.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("LEFTPADDING", (0, 0), (-1, -1), 0)]))
        flow = [ht, Spacer(1, 8)]
    else:
        flow = head + [Spacer(1, 8)]
    for si, s in enumerate(m["sections"]):
        if si:
            flow.append(PageBreak())
        flow.append(P(s["title"], "h2"))
        for n in s["notes"]:
            flow.append(P(n, "note"))
        for i in range(0, len(s["kpis"]), 4):
            row = s["kpis"][i:i + 4]
            cells = [[P(k["label"], "lab"), P(fmt(k["value"], k["kind"]), "val"), P(k["sub"], "lab")] for k in row] + [""] * (4 - len(row))
            t = Table([cells], colWidths=[W / 4] * 4)
            t.setStyle(TableStyle([("BOX", (0, 0), (-1, -1), .5, hx(LINE)), ("INNERGRID", (0, 0), (-1, -1), .5, hx(LINE)), ("VALIGN", (0, 0), (-1, -1), "TOP"),
                                   ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#F6F8FA")), ("TOPPADDING", (0, 0), (-1, -1), 6), ("BOTTOMPADDING", (0, 0), (-1, -1), 6)]))
            flow += [t, Spacer(1, 4)]
        cw = (W - 6) / 2
        grp_n = {}
        for c in s["charts"]:
            if c["kind"] == "prtg":
                grp_n[c["group"]] = grp_n.get(c["group"], 0) + 1
        last = None
        for c in [c for c in s["charts"] if c["kind"] == "prtg"]:
            if c["group"] != last:
                last = c["group"]
                flow.append(P(f"{last or 'Other'} · {grp_n[last]} link{'s' if grp_n[last] != 1 else ''}", "h3"))
            rows = [[P(h, "ch") for h in ("Channel", "Average", "Total")]] + [[P(v, "c") for v in r] for r in c["rows"]]
            tb = Table(rows, colWidths=[W * .30 * .30, W * .30 * .33, W * .30 * .37])
            tb.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), hx(NAVY)), ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F6F8FA")]),
                                    ("LINEBELOW", (0, 0), (-1, -1), .25, hx(LINE)), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3)]))
            lt = Table([[_pdf_chart(c, W * .68, 178), tb]], colWidths=[W * .70, W * .30])
            lt.setStyle(TableStyle([("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (-1, -1), 0), ("BOTTOMPADDING", (0, 0), (-1, -1), 6)]))
            flow.append(lt)
        s = {**s, "charts": [c for c in s["charts"] if c["kind"] != "prtg"]}
        for i in range(0, len(s["charts"]), 2):
            row = [_pdf_chart(c, cw, 165) for c in s["charts"][i:i + 2]]
            row += [""] * (2 - len(row))
            gt = Table([row], colWidths=[cw + 3, cw + 3])
            gt.setStyle(TableStyle([("LEFTPADDING", (0, 0), (-1, -1), 0), ("RIGHTPADDING", (0, 0), (-1, -1), 0), ("TOPPADDING", (0, 0), (-1, -1), 2), ("BOTTOMPADDING", (0, 0), (-1, -1), 2)]))
            flow.append(gt)
        for t in s["tables"]:
            ttl = P(t["title"], "h3")
            data = [[P(h, "ch") for h in t["heads"]]] + [[P(fmt(v, t["kinds"][c]), "c") for c, v in enumerate(row)] for row in t["rows"]]
            wts = [3.4 if k == "t" else 1.2 for k in t["kinds"]]
            wts[0] = max(wts[0], 2.6)
            tb = Table(data, colWidths=[W * w / sum(wts) for w in wts], repeatRows=1)
            tb.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), hx(NAVY)), ("VALIGN", (0, 0), (-1, -1), "MIDDLE"), ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#F6F8FA")]),
                                    ("LINEBELOW", (0, 0), (-1, -1), .25, hx(LINE)), ("TOPPADDING", (0, 0), (-1, -1), 3), ("BOTTOMPADDING", (0, 0), (-1, -1), 3)]))
            flow.append(KeepTogether([ttl, tb]) if len(t["rows"]) <= 14 else [ttl, tb][0])
            if len(t["rows"]) > 14:
                flow.append(tb)

    def deco(canvas, doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 7.5)
        canvas.setFillColor(hx(MUTED))
        canvas.drawString(12 * mm, 8 * mm, f"{m['title']} · {m['period']}")
        canvas.drawRightString(landscape(A4)[0] - 12 * mm, 8 * mm, f"Page {doc.page}")
        canvas.restoreState()
    out = BytesIO()
    SimpleDocTemplate(out, pagesize=landscape(A4), leftMargin=12 * mm, rightMargin=12 * mm, topMargin=12 * mm, bottomMargin=14 * mm, title=m["title"], author=m["by"]).build(flow, onFirstPage=deco, onLaterPages=deco)
    return out.getvalue()


# ---------------------------------------------------------------- PowerPoint
def to_pptx(m: dict) -> bytes:
    """Utama: dari template Pelindo (report_pptx). Bila template/pustaka bermasalah, jatuh ke layout polos supaya unduhan tetap berhasil."""
    try:
        from .report_pptx import build_from_template
        return build_from_template(m)
    except Exception:
        import logging
        logging.getLogger(__name__).exception("PPT from template failed, using the plain layout")
        return _to_pptx_plain(m)


def _to_pptx_plain(m: dict) -> bytes:
    from pptx import Presentation
    from pptx.chart.data import CategoryChartData
    from pptx.dml.color import RGBColor
    from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.util import Inches, Pt
    rgb = lambda h: RGBColor.from_string(h)
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    blank = prs.slide_layouts[6]

    def text(slide, x, y, w, h, t, size=14, bold=False, color="10171E"):
        tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.text = str(t)
        p = tf.paragraphs[0]
        p.font.size, p.font.bold, p.font.color.rgb = Pt(size), bold, rgb(color)
        return tb

    def header(slide, title, sub=""):
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(0.95))
        bar.fill.solid()
        bar.fill.fore_color.rgb = rgb(NAVY)
        bar.line.fill.background()
        text(slide, 0.5, 0.18, 9, 0.6, title, 26, True, "FFFFFF")
        text(slide, 8.3, 0.3, 4.6, 0.4, sub, 12, False, "C5D3E6")
        text(slide, 0.5, 7.05, 12, 0.3, f"{m['title']} · {m['period']} · dibuat {m['generated']}", 9, False, MUTED)

    # sampul
    s0 = prs.slides.add_slide(blank)
    bg = s0.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg.fill.solid()
    bg.fill.fore_color.rgb = rgb(NAVY)
    bg.line.fill.background()
    text(s0, 0.9, 2.4, 11.5, 1.2, m["title"], 38, True, "FFFFFF")
    text(s0, 0.9, 3.7, 11.5, 0.6, f"{m['subtitle']} · {m['period']}", 22, False, "C5D3E6")
    text(s0, 0.9, 4.5, 11.5, 0.5, " · ".join(m["filters"]), 14, False, "C5D3E6")
    text(s0, 0.9, 6.5, 11.5, 0.4, f"Generated {m['generated']} by {m['by']}", 11, False, "9FB3CC")

    for s in m["sections"]:
        sl = prs.slides.add_slide(blank)
        header(sl, s["title"], m["period"])
        y = 1.2
        for n in s["notes"]:
            text(sl, 0.5, y, 12.3, 0.4, n, 13, False, MUTED)
            y += 0.5
        kp = s["kpis"][:8]
        per = 4
        for i, k in enumerate(kp):
            col, row = i % per, i // per
            x, yy, w = 0.5 + col * 3.1, y + row * 1.35, 2.95
            card = sl.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(yy), Inches(w), Inches(1.2))
            card.fill.solid()
            card.fill.fore_color.rgb = rgb("F2F5F9")
            card.line.color.rgb = rgb(LINE)
            text(sl, x + 0.1, yy + 0.05, w - 0.2, 0.3, k["label"], 11, False, MUTED)
            text(sl, x + 0.1, yy + 0.32, w - 0.2, 0.55, fmt(k["value"], k["kind"]), 26, True, NAVY)
            text(sl, x + 0.1, yy + 0.86, w - 0.2, 0.3, k["sub"], 10, False, MUTED)
        y += (1.35 * (-(-len(kp) // per))) if kp else 0
        for ch in s["charts"]:
            cd = CategoryChartData()
            cd.categories = ch["labels"]
            for se in ch["series"]:
                cd.add_series(se["name"], [v or 0 for v in se["values"]])
            h = max(2.2, 7.0 - y - 0.1)
            gf = sl.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.5), Inches(y), Inches(12.3), Inches(h), cd)
            c = gf.chart
            c.has_title = True
            c.chart_title.text_frame.text = ch["title"]
            c.chart_title.text_frame.paragraphs[0].font.size = Pt(13)
            c.has_legend = len(ch["series"]) > 1
            if c.has_legend:
                c.legend.position, c.legend.include_in_layout = XL_LEGEND_POSITION.BOTTOM, False
            c.category_axis.tick_labels.font.size = Pt(8)
            c.value_axis.tick_labels.font.size = Pt(9)
            for j, se in enumerate(c.plots[0].series):
                se.format.fill.solid()
                se.format.fill.fore_color.rgb = rgb(PAL[j % len(PAL)])
            break  # satu grafik per slide ringkas; grafik lain ada di Excel
        for t in s["tables"]:
            for start in range(0, len(t["rows"]), 12):
                rows = t["rows"][start:start + 12]
                ts = prs.slides.add_slide(blank)
                header(ts, s["title"], t["title"] + (f" ({start // 12 + 1})" if len(t["rows"]) > 12 else ""))
                nr, nc = len(rows) + 1, len(t["heads"])
                shape = ts.shapes.add_table(nr, nc, Inches(0.5), Inches(1.2), Inches(12.3), Inches(0.4 * nr))
                tbl = shape.table
                wts = [3.4 if k == "t" else 1.2 for k in t["kinds"]]
                wts[0] = max(wts[0], 2.6)
                for c in range(nc):
                    tbl.columns[c].width = Inches(12.3 * wts[c] / sum(wts))
                for c, h in enumerate(t["heads"]):
                    cell = tbl.cell(0, c)
                    cell.text = str(h)
                    cell.text_frame.paragraphs[0].font.size, cell.text_frame.paragraphs[0].font.bold = Pt(11), True
                for r, row in enumerate(rows, 1):
                    for c, v in enumerate(row):
                        cell = tbl.cell(r, c)
                        cell.text = fmt(v, t["kinds"][c])
                        cell.text_frame.paragraphs[0].font.size = Pt(10)
    out = BytesIO()
    prs.save(out)
    return out.getvalue()


RENDER = {"xlsx": to_xlsx, "pdf": to_pdf, "pptx": to_pptx}
