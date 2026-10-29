"""PowerPoint dari template Pelindo (backend/assets/template_pelindo.pptx): sampul, header gradien, strip KPI, grafik asli, dan slide penutup
dipakai dari template; isi diganti dengan angka laporan. Slide template yang berisi grafik contoh dibuang."""
import copy
from io import BytesIO
from pathlib import Path
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Pt
from .report import fmt

TEMPLATE = Path(__file__).resolve().parent.parent / "assets" / "template_pelindo.pptx"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
NAVY, BLUE, CYAN, GRAY, LINE, ZEBRA = "0B2F6B", "0B75BB", "25B7E5", "6B7280", "D9E2EC", "F3F6FA"
BAND = {"servicedesk": NAVY, "noc": BLUE, "soc": CYAN}
PAL = ["4F81BD", "9BBB59", "4BACC6", "2C4D75", "F79646", "8064A2", "C0504D", "7F7F7F"]
CHART = {"bar": XL_CHART_TYPE.COLUMN_CLUSTERED, "stack": XL_CHART_TYPE.COLUMN_STACKED, "hbar": XL_CHART_TYPE.BAR_CLUSTERED,
         "donut": XL_CHART_TYPE.DOUGHNUT, "line": XL_CHART_TYPE.LINE, "prtg": XL_CHART_TYPE.LINE, "area": XL_CHART_TYPE.AREA}
X0, W0, Y0, Y1 = 40, 1361, 101, 735  # area isi (pt): sama dengan tabel KPI di template


def rgb(h):
    return RGBColor.from_string(h)


def _named(slide, key):
    for s in slide.shapes:
        if f";{key};" in s.name:
            return s
    return None


def _set_runs(par, text):
    runs = par.runs
    if runs:
        runs[0].text = text
        for r in runs[1:]:
            r._r.getparent().remove(r._r)
    else:
        par.text = text


def _set_text(shape, text):
    if shape is None:
        return
    tf = shape.text_frame
    _set_runs(tf.paragraphs[0], text)
    for p in tf.paragraphs[1:]:
        p._p.getparent().remove(p._p)


def _clone(src, dst):
    """Salin bentuk dekoratif (header gradien, logo, pola, teks periode) dari slide template; grafik/tabel contoh dilewati."""
    for shp in src.shapes:
        if shp._element.tag.endswith("graphicFrame"):
            continue
        el = copy.deepcopy(shp._element)
        for node in el.iter():
            for attr in ("embed", "link", "id"):
                key = f"{{{R_NS}}}{attr}"
                old = node.get(key)
                if old and old in src.part.rels and not src.part.rels[old].is_external:
                    rel = src.part.rels[old]
                    node.set(key, dst.part.relate_to(rel.target_part, rel.reltype))
        dst.shapes._spTree.append(el)


def _style(p, size, bold, color):
    """Gaya ke paragraf dan setiap run supaya konsisten di PowerPoint maupun LibreOffice."""
    for font in [p.font] + [r.font for r in p.runs]:
        font.name, font.size, font.bold, font.color.rgb = "Arial", Pt(size), bold, rgb(color)


def _tb(slide, x, y, w, h, text, size=14, bold=False, color=GRAY, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.TOP):
    box = slide.shapes.add_textbox(Pt(x), Pt(y), Pt(w), Pt(h))
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Pt(4)
    tf.text = str(text)
    p = tf.paragraphs[0]
    p.alignment = align
    _style(p, size, bold, color)
    return box


def _rect(slide, x, y, w, h, color):
    r = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Pt(x), Pt(y), Pt(w), Pt(h))
    r.fill.solid()
    r.fill.fore_color.rgb = rgb(color)
    r.line.fill.background()
    r.shadow.inherit = False
    return r


def _kpi_strip(slide, y, groups, ncols, band=True):
    """groups = [(judul pita, warna, [kpi...])]; tiap KPI satu kolom. Gaya sama dengan tabel KPI di template. band=False: hanya sel KPI."""
    cw = W0 / ncols
    col = 0
    top = y + (70 if band else 0)
    for title, color, kpis in groups:
        if band:
            b = _rect(slide, X0 + col * cw, y, cw * len(kpis), 62, color)
            tf = b.text_frame
            tf.text = title.upper()
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.CENTER
            _style(p, 20, True, "FFFFFF")
        for k in kpis:
            x = X0 + col * cw
            val = fmt(k["value"], k["kind"])
            lab = 15 if cw > 200 else 12 if cw > 140 else 10.5                       # label and value shrink with the column so neighbouring figures never overlap
            vs = max(12, min(26, int(cw / (max(len(val), 1) * 0.64))))
            _tb(slide, x, top, cw, 40, k["label"].upper(), lab, True, GRAY, anchor=MSO_ANCHOR.BOTTOM)
            _tb(slide, x, top + 42, cw, 40, val, vs, True, NAVY, anchor=MSO_ANCHOR.MIDDLE)
            _tb(slide, x + 4, top + 86, cw - 8, 34, k["sub"], 11 if cw < 140 else 12, False, GRAY)
            col += 1
    h = 124 + (70 if band else 0)
    ln = slide.shapes.add_connector(1, Pt(X0), Pt(y + h), Pt(X0 + W0), Pt(y + h))
    ln.line.color.rgb = rgb(LINE)
    return y + h + 8


def _chart(slide, ch, x, y, w, h, fs=11):
    cd = CategoryChartData()
    cd.categories = ch["labels"]
    for se in ch["series"]:
        cd.add_series(se["name"], [v or 0 for v in se["values"]])
    kind = ch.get("kind", "bar")
    c = slide.shapes.add_chart(CHART.get(kind, XL_CHART_TYPE.COLUMN_CLUSTERED), Pt(x), Pt(y), Pt(w), Pt(h), cd).chart
    c.has_title = True
    c.chart_title.text_frame.text = ch["title"]
    c.chart_title.text_frame.paragraphs[0].font.size = Pt(fs + 2)
    c.chart_title.text_frame.paragraphs[0].font.bold = True
    multi = len(ch["series"]) > 1
    c.has_legend = multi or kind == "donut"
    if c.has_legend:
        c.legend.position = XL_LEGEND_POSITION.RIGHT if kind == "donut" else XL_LEGEND_POSITION.BOTTOM
        c.legend.include_in_layout = False
        c.legend.font.size = Pt(fs - 1)
    plot = c.plots[0]
    if kind == "donut":
        plot.has_data_labels = True
        plot.data_labels.font.size = Pt(fs - 1)
        plot.data_labels.show_value = True
        for i in range(len(ch["labels"])):
            pt = plot.series[0].points[i]
            pt.format.fill.solid()
            pt.format.fill.fore_color.rgb = rgb(PAL[i % len(PAL)])
        return c
    c.category_axis.tick_labels.font.size = Pt(fs - 1)
    c.value_axis.tick_labels.font.size = Pt(fs - 1)
    c.value_axis.major_gridlines.format.line.color.rgb = rgb(LINE)
    if kind == "hbar":
        c.category_axis.reverse_order = True
    for j, se in enumerate(plot.series):
        if kind in ("line", "prtg"):
            se.format.line.color.rgb = rgb(ch["series"][j].get("color") or PAL[j % len(PAL)])
            se.smooth = False
        else:
            se.format.fill.solid()
            se.format.fill.fore_color.rgb = rgb(ch["series"][j].get("color") or PAL[j % len(PAL)])
            if kind == "area":                                          # see-through fill so the two directions can overlap, as in PRTG graphs
                try:
                    from pptx.oxml.ns import qn
                    from lxml import etree
                    etree.SubElement(se._element.find(".//" + qn("a:srgbClr")), qn("a:alpha")).set("val", "55000")
                except Exception:
                    pass
    return c


def _table(slide, t, rows, y):
    nr, nc = len(rows) + 1, len(t["heads"])
    shape = slide.shapes.add_table(nr, nc, Pt(X0), Pt(y), Pt(W0), Pt(30 * nr))
    tbl = shape.table
    wts = [3.4 if k == "t" else 1.2 for k in t["kinds"]]
    wts[0] = max(wts[0], 2.4)
    for c in range(nc):
        tbl.columns[c].width = Pt(W0 * wts[c] / sum(wts))
    for r in range(nr):
        tbl.rows[r].height = Pt(30)

    def put(r, c, text, size, bold, color, fill, align):
        cell = tbl.cell(r, c)
        cell.fill.solid()
        cell.fill.fore_color.rgb = rgb(fill)
        cell.vertical_anchor = MSO_ANCHOR.MIDDLE
        cell.text = text
        p = cell.text_frame.paragraphs[0]
        p.alignment = align
        _style(p, size, bold, color)
    for c, h in enumerate(t["heads"]):
        put(0, c, str(h), 14, True, "FFFFFF", NAVY, PP_ALIGN.LEFT if t["kinds"][c] == "t" else PP_ALIGN.RIGHT)
    for r, row in enumerate(rows, 1):
        for c, v in enumerate(row):
            put(r, c, fmt(v, t["kinds"][c]), 13, False, "10171E", ZEBRA if r % 2 == 0 else "FFFFFF", PP_ALIGN.LEFT if t["kinds"][c] == "t" else PP_ALIGN.RIGHT)


def _find_slide(slides, *keys):
    return next((s for s in slides if all(_named(s, k) is not None for k in keys)), None)


def build_from_template(m: dict) -> bytes:
    prs = Presentation(str(TEMPLATE))
    sl = list(prs.slides)
    cover = _find_slide(sl, "166", "167")
    content_src = _find_slide(sl[1:], "190", "194") or sl[1]
    thanks = sl[-1]
    secs = [s for s in m["sections"] if s["key"] != "ringkasan"]
    ov = next((s for s in m["sections"] if s["key"] == "ringkasan"), None)

    # --- sampul: judul jenis laporan, perusahaan, subjudul, periode
    t = _named(cover, "166")
    if t is not None:
        ps = t.text_frame.paragraphs
        _set_runs(ps[0], m["subtitle"])
        if len(ps) > 1:
            _set_runs(ps[1], m.get("company") or "PT PELINDO MULTI TERMINAL")
    _set_text(_named(cover, "167"), m.get("tagline") or " · ".join(m["filters"]))
    _set_text(_named(cover, "168"), m["period"])

    # --- slide isi (header, logo, dan pola dari template)
    blank = content_src.slide_layout
    new = []

    def page(title):
        s = prs.slides.add_slide(blank)
        _clone(content_src, s)
        _set_text(_named(s, "194"), title)
        _set_text(_named(s, "168"), m["period"])
        new.append(s)
        return s

    def tables_of(sec, pre):
        for t in sec["tables"]:
            for start in range(0, len(t["rows"]), 14):
                part = t["rows"][start:start + 14]
                ts = page(f"{pre} · {t['title']}" + (f" ({start // 14 + 1}/{-(-len(t['rows']) // 14)})" if len(t["rows"]) > 14 else ""))
                _table(ts, t, part, 110)

    def chart_pages(sec, pre, skip):
        """Grafik sisa dalam grid 2 x 2 per slide."""
        rest = sec["charts"][skip:]
        gw, gh = W0 / 2 - 6, (Y1 - Y0) / 2 - 6
        for i in range(0, len(rest), 4):
            s = page(f"{pre} · Charts" + (f" ({i // 4 + 1}/{-(-len(rest) // 4)})" if len(rest) > 4 else ""))
            for j, ch in enumerate(rest[i:i + 4]):
                _chart(s, ch, X0 + (j % 2) * (gw + 12), Y0 + (j // 2) * (gh + 12), gw, gh, fs=10)

    is_summary = m.get("kind") == "summary"
    if ov and ov["kpis"] and (is_summary or len(secs) > 1):
        s = page("Summary · Overview")
        groups, order = {}, []
        for k in ov["kpis"]:
            key = k.get("sec", "")
            if key not in groups:
                groups[key] = (k.get("sec_title", key), BAND.get(key, NAVY), [])
                order.append(key)
            groups[key][2].append({**k, "label": k["label"].split(" · ", 1)[-1]})
        gl = [groups[k] for k in order]
        y = _kpi_strip(s, Y0, gl, max(sum(len(g[2]) for g in gl), 1))
        charts = ov["charts"][:3]
        for i, ch in enumerate(charts):
            w = W0 / len(charts)
            _chart(s, ch, X0 + i * w, y, w - 8, Y1 - y)
        if ov["tables"]:
            tables_of(ov, "Summary")
    for sec in secs:
        pre = sec.get("prefix") or sec["title"]
        s = page(f"{pre} · Overview")
        y = Y0
        for n in sec["notes"]:
            _tb(s, X0, y, W0, 30, n, 16, False, GRAY, PP_ALIGN.LEFT)
            y += 36
        used = 0
        if sec["kpis"]:
            col = BAND.get(sec["key"], NAVY)
            y = _kpi_strip(s, y, [(sec["title"], col, sec["kpis"][:6])], min(len(sec["kpis"]), 6))
            if len(sec["kpis"]) > 6:
                y = _kpi_strip(s, y, [("", col, sec["kpis"][6:12])], min(len(sec["kpis"]) - 6, 6), band=False)
        if sec["charts"]:
            first = sec["charts"][:2]
            wts = [0.42, 0.58] if len(first) == 2 else [1]
            x = X0
            for ch, wt in zip(first, wts):
                _chart(s, ch, x, y, W0 * wt - 8, max(180, Y1 - y))
                x += W0 * wt
            used = len(first)
        chart_pages(sec, pre, used)
        tables_of(sec, pre)

    # --- penutup (slide Thank You dari template): kontak dari parameter
    c = m.get("contact") or {}
    _set_text(_named(thanks, "451"), c.get("phone") or "")
    _set_text(_named(thanks, "453"), c.get("email") or "")
    addr = _named(thanks, "456")
    if addr is not None and c.get("address"):
        name, _, rest = str(c["address"]).partition(",")
        ps = addr.text_frame.paragraphs
        _set_runs(ps[0], name.strip())
        if len(ps) > 1:
            _set_runs(ps[1], rest.strip())

    # --- urutan: sampul, isi baru, penutup; slide contoh template dibuang
    lst = prs.slides._sldIdLst
    ids = list(lst)
    n_old = len(sl)
    first_id, last_id = ids[sl.index(cover)], ids[n_old - 1]
    for sid in ids[:n_old]:
        if sid is not first_id and sid is not last_id:
            prs.part.drop_rel(sid.rId)
    new_ids = ids[n_old:]
    for sid in ids:
        lst.remove(sid)
    for sid in [first_id] + new_ids + [last_id]:
        lst.append(sid)
    for i, s in enumerate(prs.slides, 1):
        for key in ("164", "193", "457"):
            _set_text(_named(s, key), str(i))
    out = BytesIO()
    prs.save(out)
    return out.getvalue()
